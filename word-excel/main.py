import pandas as pd
import xlwt
from docx import Document

from config import PATHS
from util import remove


word = PATHS['input_path'] + '/' + '2022年照明服务所应急防汛值班表.docx'
_excel = PATHS['output_path'] + '/' + '_值班管理.xlsx'
excel = PATHS['output_path'] + '/' + '值班管理.xlsx'

doc = Document(word)
xls = pd.ExcelWriter(_excel)

tables = doc.tables  # word 中所有 tables

mat = []
for n in range(0,len(tables)):
    tb = tables[n]
    print(len(tb.rows), len(tb.columns),"第{0}个表格".format(n))  # 行数、列数

    for i in range(2,len(tb.rows)):# 读取行
        row = []
        for j in range(0,len(tb.columns)):# 读取列
            cell = tb.cell(i,j)# 读取单元格
            txt = cell.text if cell.text != "" else " " # 无内容用空格占位
            row.append(txt)
        mat.append(row)

workbook = xlwt.Workbook(encoding = 'utf-8') # 创建工作簿
xlsheet = workbook.add_sheet("Sheet1",cell_overwrite_ok=True)

# 添加表头
table_head = ['值班日期','带班人','带班人电话','成员','成员电话']
headlen = len(table_head)
for i in range(headlen):
    xlsheet.write(0,i,table_head[i])

# 写入数据
for i in range(len(mat)):
    for j in range(len(row)):
        xlsheet.write(i+1,j,mat[i][j])
workbook.save(_excel)

print("Done!------------")